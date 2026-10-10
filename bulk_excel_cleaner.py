
from pathlib import Path
from datetime import datetime
import csv
import re
import tkinter as tk
from tkinter import filedialog, messagebox

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

VERSION = "7.1"

COLUMN_ALIASES = {
    "name": "Name",
    "full name": "Name",
    "customer name": "Name",
    "email": "Email",
    "e-mail": "Email",
    "email address": "Email",
    "e-mail address": "Email",
    "city": "City",
    "town": "City",
    "phone": "Phone",
    "mobile": "Phone",
    "mobile number": "Phone",
    "phone number": "Phone",
    "contact number": "Phone",
}

EMAIL_PATTERN = re.compile(
    r"^[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}$"
)


def normalize_header(header):
    text = str(header or "").strip()
    key = " ".join(text.lower().split())
    return COLUMN_ALIASES.get(key, text)


def normalize_phone(value):
    if value is None:
        return ""

    phone = str(value).strip()

    if phone.endswith(".0"):
        phone = phone[:-2]

    phone = re.sub(r"[\s()-]", "", phone)

    if phone.startswith("+91"):
        phone = phone[3:]
    elif phone.startswith("91") and len(phone) == 12:
        phone = phone[2:]

    return phone


def valid_phone(phone):
    return bool(re.fullmatch(r"[6-9]\d{9}", phone))


def read_excel(path):
    book = load_workbook(
        path,
        read_only=True,
        data_only=True
    )

    try:
        sheet = book.active
        for row in sheet.iter_rows(values_only=True):
            yield row
    finally:
        book.close()


def read_csv(path):
    with path.open(
        "r",
        encoding="utf-8-sig",
        newline=""
    ) as file:
        yield from csv.reader(file)


def format_sheet(sheet):
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions

    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill(
            "solid",
            fgColor="17365D"
        )
        cell.alignment = Alignment(
            horizontal="center"
        )

    for row in sheet.iter_rows(min_row=2):
        if row[0].row % 2 == 0:
            for cell in row:
                cell.fill = PatternFill(
                    "solid",
                    fgColor="EAF2F8"
                )

    for column in sheet.columns:
        max_length = max(
            len(str(cell.value))
            if cell.value is not None else 0
            for cell in column
        )

        letter = get_column_letter(
            column[0].column
        )

        sheet.column_dimensions[letter].width = min(
            max(max_length + 4, 14),
            45
        )


def main():
    root = tk.Tk()
    root.withdraw()

    selected = filedialog.askopenfilenames(
        title="Select Excel or CSV files to clean",
        filetypes=[
            ("Excel and CSV", "*.xlsx *.csv"),
            ("Excel Files", "*.xlsx"),
            ("CSV Files", "*.csv"),
        ]
    )

    if not selected:
        root.destroy()
        return

    # Let user choose output location
    output_filename = filedialog.asksaveasfilename(
        title="Save Cleaned Excel Report",
        defaultextension=".xlsx",
        initialfile="bulk_cleaned_customers.xlsx",
        filetypes=[
            ("Excel Workbook", "*.xlsx")
        ]
    )

    if not output_filename:
        root.destroy()
        return

    output_file = Path(output_filename)

    # Prevent overwriting input files
    selected_paths = {
        Path(filename).resolve()
        for filename in selected
    }

    if output_file.resolve() in selected_paths:
        messagebox.showerror(
            "Invalid Save Location",
            "Please choose a different filename from the input files."
        )
        root.destroy()
        return

    workbook = Workbook()

    clean = workbook.active
    clean.title = "Clean Data"

    missing = workbook.create_sheet("Missing Emails")
    invalid_email = workbook.create_sheet("Invalid Emails")
    invalid_phone = workbook.create_sheet("Invalid Phones")
    summary = workbook.create_sheet("Summary")
    errors = workbook.create_sheet("File Errors")

    errors.append(["File", "Problem"])

    headers = None
    seen = set()

    stats = {
        "Files Processed": 0,
        "Total Records": 0,
        "Duplicates Removed": 0,
        "Missing Emails": 0,
        "Invalid Emails": 0,
        "Invalid Phones": 0,
        "Clean Customers": 0,
        "Files Skipped": 0,
    }

    for filename in selected:
        path = Path(filename)

        if path.name.startswith("~$"):
            continue

        try:
            if path.suffix.lower() == ".csv":
                rows = read_csv(path)
            else:
                rows = read_excel(path)

            try:
                original_headers = next(rows, None)

                if not original_headers:
                    raise ValueError("File is empty")

                file_headers = [
                    normalize_header(h)
                    for h in original_headers
                ]

                if len(set(file_headers)) != len(file_headers):
                    raise ValueError(
                        "Duplicate column names after normalization"
                    )

                if "Email" not in file_headers:
                    raise ValueError("Email column missing")

                if headers is None:
                    headers = file_headers

                    for sheet in [
                        clean,
                        missing,
                        invalid_email,
                        invalid_phone
                    ]:
                        sheet.append(headers)

                if set(file_headers) != set(headers):
                    raise ValueError(
                        "Columns do not match the first file"
                    )

                stats["Files Processed"] += 1

                for row in rows:
                    if not row or all(
                        value is None or str(value).strip() == ""
                        for value in row
                    ):
                        continue

                    if len(row) > len(file_headers):
                        raise ValueError(
                            "Row has more values than headers"
                        )

                    row_data = dict(zip(file_headers, row))

                    email = row_data.get("Email")
                    email = (
                        str(email).strip().lower()
                        if email is not None
                        else ""
                    )
                    row_data["Email"] = email

                    if "Phone" in headers:
                        row_data["Phone"] = normalize_phone(
                            row_data.get("Phone")
                        )

                    record = tuple(
                        row_data.get(column)
                        for column in headers
                    )

                    stats["Total Records"] += 1

                    if record in seen:
                        stats["Duplicates Removed"] += 1
                        continue

                    seen.add(record)

                    if not email:
                        missing.append(record)
                        stats["Missing Emails"] += 1

                    elif not EMAIL_PATTERN.fullmatch(email):
                        invalid_email.append(record)
                        stats["Invalid Emails"] += 1

                    elif (
                        "Phone" in headers
                        and not valid_phone(
                            row_data.get("Phone", "")
                        )
                    ):
                        invalid_phone.append(record)
                        stats["Invalid Phones"] += 1

                    else:
                        clean.append(record)
                        stats["Clean Customers"] += 1

            finally:
                rows.close()

        except (
            OSError,
            ValueError,
            UnicodeError,
            KeyError,
            TypeError
        ) as error:
            stats["Files Skipped"] += 1
            errors.append([path.name, str(error)])
            print("Problem with", path.name, ":", error)

    summary.append(["Metric", "Value"])
    summary.append(["Tool Version", VERSION])
    summary.append([
        "Generated At",
        datetime.now().strftime("%Y-%m-%d %H:%M")
    ])

    for name, value in stats.items():
        summary.append([name, value])

    for sheet in workbook.worksheets:
        format_sheet(sheet)

    try:
        workbook.save(output_file)
    except (PermissionError, OSError) as error:
        messagebox.showerror(
            "Save Failed",
            f"Could not save the report:\n{error}"
        )
        root.destroy()
        return

    result = "\n".join(
        f"{name}: {value}"
        for name, value in stats.items()
    )

    print("Excel Data Cleaner 7.1 completed!")
    print(result)
    print("Output:", output_file)

    messagebox.showinfo(
        "Cleaning Complete",
        f"Excel Data Cleaner {VERSION} finished!\n\n"
        f"{result}\n\n"
        f"Saved to:\n{output_file}"
    )

    root.destroy()


if __name__ == "__main__":
    main()
